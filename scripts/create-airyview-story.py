from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "share"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "AiryView_AIとの会話から生まれたPDFアプリ.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.72)
sec.left_margin, sec.right_margin = Inches(.85), Inches(.85)

styles = doc.styles
styles["Normal"].font.name = "Yu Gothic"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.35
for name, size in [("Title", 27), ("Heading 1", 19), ("Heading 2", 13)]:
    s = styles[name]
    s.font.name = "Yu Gothic"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(14)
    s.paragraph_format.space_after = Pt(8)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def margins(cell, top=110, start=150, bottom=110, end=150):
    tc = cell._tc.get_or_add_tcPr()
    m = tc.first_child_found_in("w:tcMar")
    if m is None:
        m = OxmlElement("w:tcMar"); tc.append(m)
    for side, val in [("top",top),("start",start),("bottom",bottom),("end",end)]:
        el=OxmlElement("w:"+side); el.set(qn("w:w"),str(val)); el.set(qn("w:type"),"dxa"); m.append(el)

def conversation(who, text, dark=False):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.15)
    table.columns[1].width = Inches(5.4)
    row = table.rows[0]
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    left, right = row.cells
    left.width, right.width = Inches(1.15), Inches(5.4)
    shade(left, "202020" if dark else "E7E9EC")
    shade(right, "F3F4F6" if dark else "FFFFFF")
    for c in (left,right): margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=left.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(who); r.bold=True; r.font.color.rgb=RGBColor(255,255,255) if dark else RGBColor(0,0,0)
    right.paragraphs[0].add_run(text)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

# Cover
p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run("AIとの会話から生まれた仕事のアプリ")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("AiryViewを一緒に作った記録"); r.bold=True; r.font.size=Pt(16)
icon=ROOT/"assets"/"icons"/"AiryView-symbol-preview.png"
if icon.exists():
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(icon), width=Inches(2.0))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("私はプログラムコードを一度も書いたことがありません。\nそれでもAIとの会話を通じて、仕事用のアプリやプラグインをいくつか作ってきました。\nこれまではClaudeを使い、最近はChatGPTのCodexを中心に開発しています。AiryViewは、9月13日の朝にPDFリーダーの枠組みができ、その後、細かな改善やメモ機能などを追加しました。")
r.font.size=Pt(13)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before=Pt(22)
p.add_run("2026年9月  制作記録").italic=True

doc.add_page_break()
doc.add_heading("始まりは仕事の困りごと", level=1)
doc.add_paragraph("私は、仕事で使う1/5000の図面を正確に出力したいと考えていました。Adobe Acrobat Readerでは、1/5000の図面を原寸どおりに印刷でき、図面上の1cmが現地の50mを表すスケールも正確に合っていました。ただ、私の環境では立ち上がりが重く感じられ、案内や広告も多いことが気になっていました。反対に、ブラウザは立ち上がりが速く、軽くて使いやすいのですが、原寸どおりに出力する設定を選んでも、印刷した紙の1cmが正確に合わないことが何度もありました。そこで、Adobeのように寸法を正確に出力でき、ブラウザのように軽くて簡単な、自分向けのオリジナルなアプリが欲しいと思いました。そこでAIに、仕様書ではなく普段の言葉で相談しました。")
doc.add_heading("話しているうちに形になった", level=2)
doc.add_paragraph("9月13日の朝に、『寸法を正確に印刷できて、軽くて使いやすいPDFリーダーが欲しい』とAIへ話しました。すると、その日の朝のうちに基本的な形ができ、実際にPDFを開いて試せるようになりました。私はコードを書かず、できたものを触って、気づいたことを普段の言葉で返しただけです。その後も同じやり取りを重ね、細かな使い勝手やメモ機能などを追加していきました。私が実際にAIと会話していた時間は、肌感覚では合計2〜3時間ほどです。特別な知識がなくても、欲しいものを話し、試した感想を伝えていくことで、使えるアプリが短い時間で形になりました。")
doc.add_paragraph("昨夜は晩酌をしながら、思いついた機能を会話で少しずつ追加していました。まとまった開発時間を確保したというより、普段の時間の中で楽しみながら形にしていった、という感覚です。")
doc.add_page_break()
conversation("私", "Adobe Acrobat Readerと同じくらい、寸法に正確に出力できるPDFリーダーを作りたいです。", True)
conversation("AI", "まず、原寸や指定倍率では勝手に縮小せず、用紙からはみ出す場合はプレビューで分かるようにしましょう。")
doc.add_paragraph("最初の試作を実際に触ると、PDFを開いた直後の大きさや、ページの移動方法にも希望が出てきました。")
conversation("私", "PDFは最初から100％で表示してください。\n何ページもある資料は、縦にスクロールして読みたいです。", True)
doc.add_paragraph("私はプログラムコードを一度も書いたことがありません。一方で、これまでのアプリ作りを通じて、仕事のどこで困っているか、実際に使うと何が邪魔になるか、結果が正しいかを確かめる経験は積んできました。AiryViewでも最初から全体を決め切らず、できた画面を触り、気づいたことをそのまま話していきました。")


doc.add_heading("これまでにも少し作っていた", level=1)
doc.add_paragraph("AiryViewが最初のアプリだったわけではありません。これまでにもAIとの会話で、仕事用のアプリやプラグインをいくつか作ってみました。コードは書かず、欲しい動きや使った感想を言葉で伝え、できたものを試しながら直してきました。")
doc.add_paragraph("これまではClaudeを使うことが多く、最近はChatGPTのCodexを中心に開発しています。使うAIは変わっても、普段の言葉で相談し、実際に触って感想を返す進め方は同じです。AiryViewでも『原寸で出したい』『100％で止めたい』『印刷ボタンが隠れる』という短い言葉から形になっていきました。")

doc.add_heading("画面を見ながら少しずつ整える", level=1)
conversation("私", "PDFを画面の横幅いっぱいに表示するボタンは、意味が分かる名前にしてください。回転は文字ではなくアイコンにして、上のボタンもシンプルにしてください。", True)
conversation("AI", "開く、印刷、倍率、回転、画面幅の操作を中心に整理します。ショートカットも表示します。")
conversation("私", "倍率は手入力できるように。Ctrlとスクロールでは必ず100％に止まってほしいです。", True)
conversation("AI", "倍率をまたいだときに100％で一度止まるようにします。プラスとマイナスも同じ動きに揃えます。")
doc.add_paragraph("私が伝えたのは、『何となく使いづらい』『ここが隠れる』『前の色の方がよい』といった感想です。スクリーンショットを見せると、AIが修正版を作り、私はもう一度使って確かめました。")

doc.add_heading("実際の仕事で試して完成度を上げる", level=1)
conversation("私", "会社の複合機で印刷を確認します。", True)
conversation("AI", "原本の用紙サイズと向きを引き継ぎ、印刷倍率100％で確認できるようにします。")
conversation("私", "会社で印刷しました。0.1mmくらいずれて見えますが、私の用途では完璧です。", True)
doc.add_paragraph("画面上で動くだけでは、仕事のアプリとして十分か分かりません。実際の複合機で印刷し、自分の用途に必要な精度を満たすか確認したことで、AiryViewは『試作品』から『使えるアプリ』へ進みました。")

doc.add_heading("会話は見た目や名前にも及んだ", level=1)
conversation("私", "黒背景で、グラデーションなし。アイコンは少しかわいくしたいです。", True)
conversation("AI", "羽根、勾玉、紙などを試し、最後に折り紙のツルをシンボルとして提案します。")
conversation("私", "ツル、採用。左上やタスクバーでもはっきり見えるようにしてください。", True)
conversation("AI", "小さい表示専用に首と面を太くし、アプリ、タスクバー、エクスプローラーで揃えます。")
doc.add_paragraph("機能だけでなく、色、言葉、ボタンの位置、アイコンまで会話の中で変わりました。私にも好みや使い勝手は判断できたので、それを伝え、AIが形にする往復が続きました。")

doc.add_heading("やってみて感じたこと", level=1)
items = [
    ("専門用語をほとんど使わなかった", "私が使ったのは、『ボタンが隠れる』『少しぼやける』『閉じるのがひと手間』といった普段の言葉でした。"),
    ("途中で考えが変わった", "色やアイコンは何度か試し、前の案へ戻したこともあります。それでも作業は続けられました。"),
    ("画像が会話の助けになった", "説明しにくいときは、画面のスクリーンショットや対象のPDF、プリンター名を見せました。"),
    ("コードを書かなくても作り続けられた", "技術的な方法はAIに任せ、私は使い方や、実際に試した結果の良し悪しを伝えました。複数のアプリで同じ進め方を重ねています。"),
    ("実際に使うことが判断材料になった", "会社の複合機で印刷し、自分の仕事に十分な精度だと確認できたことで、ひとつの区切りがつきました。"),
]
for title, body in items:
    p=doc.add_paragraph(); r=p.add_run(title+"  "); r.bold=True; p.add_run(body)

doc.add_heading("私の場合はこんな始まり方だった", level=1)
doc.add_paragraph("立派な企画書や完成図があったわけではありません。過去のアプリ作りで身についたのは、プログラムの書き方ではなく、仕事で困っていることを具体的に伝え、出てきたものを実際に試す進め方でした。AiryViewも、使った感想を返すたびに少しずつ変わっていきました。")
conversation("最初の相談", "仕事で使うPDFを、寸法どおりに印刷できるアプリが欲しいです。", True)

doc.add_heading("現在のAiryView", level=1)
doc.add_paragraph("AiryViewは、PDFを軽く開いて読み、原寸を意識して印刷できるWindowsアプリになりました。複数ページも自然にスクロールでき、記入や検索など、仕事でPDFを扱うための機能も備えています。さらに、Markdownや画像、テキストも同じ画面のタブで開けます。何も開かずに起動すれば、シンプルなメモ帳としても使えます。最初は自分の印刷の困りごとから始まりましたが、毎日のファイル確認やちょっとしたメモにも使えるアプリへ育ちました。")
doc.add_paragraph("私はコードを一度も書いたことがありません。ただ、AIとの会話で仕事用のアプリやプラグインをいくつか作り、実際に試して直してきました。その経験とAIの作業速度が合わさり、AiryViewの中心となるPDF閲覧・原寸印刷は短い会話の積み重ねで仕事に使えるところまで進み、その後も実際に使って見つけた問題を会話で直し続けています。")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("これはチャットの逐語録ではなく、実際の開発経過をもとに読みやすく再構成した記録です。")
r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(90,90,90)

# footer
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AiryView  AIとの会話から生まれた仕事のアプリ").font.size=Pt(8)

doc.core_properties.title = "AIとの会話から生まれた仕事のアプリ"
doc.core_properties.subject = "AiryView制作記録"
doc.save(DOCX)

check = Document(DOCX)
for checked_section in check.sections:
    width_mm = checked_section.page_width.mm
    height_mm = checked_section.page_height.mm
    if abs(width_mm - 210) > 0.1 or abs(height_mm - 297) > 0.1:
        raise RuntimeError(f"A4ではないページ設定です: {width_mm:.1f} x {height_mm:.1f} mm")

print(DOCX)
